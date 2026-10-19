import os
import re
import json
import glob
from pathlib import Path
import openpyxl

def parse_options(options_raw, options_vi_raw=None):
    """
    Tách chuỗi các lựa chọn (A) ..., B) ...) thành danh sách dictionary {key, text, text_vi}.
    Hỗ trợ ghép với chuỗi options_vi_raw (nếu có cột dịch riêng) hoặc bóc tách dịch nội tuyến.
    """
    def _extract_opts(raw_text):
        if not raw_text:
            return {}
        lines = [line.strip() for line in str(raw_text).split('\n') if line.strip()]
        opts = {}
        current_k = None
        current_t = []

        for line in lines:
            match = re.match(r'^([A-E])\)\s*(.*)', line)
            if match:
                if current_k:
                    opts[current_k] = " ".join(current_t).strip()
                current_k = match.group(1)
                current_t = [match.group(2)]
            else:
                if current_k:
                    current_t.append(line)
        if current_k:
            opts[current_k] = " ".join(current_t).strip()
        return opts

    en_opts = _extract_opts(options_raw)
    vi_opts = _extract_opts(options_vi_raw) if options_vi_raw else {}

    options = []
    for key in sorted(en_opts.keys()):
        text_en = en_opts[key]
        text_vi = vi_opts.get(key, "")

        # Kiểm tra nếu bản dịch tiếng Việt được nhúng ngay trong text EN theo dạng:
        # "English text [Dịch: Tiếng Việt]" hoặc "English text /// Tiếng Việt"
        if not text_vi:
            split_inline = re.split(r'(?:\s*\[(?:Dịch|VI|Tiếng Việt):\s*|\s*///\s*|\s*\n(?:Dịch|VI):\s*)(.*)', text_en, flags=re.IGNORECASE)
            if len(split_inline) > 1:
                text_en = split_inline[0].strip()
                text_vi = split_inline[1].rstrip(']').strip()

        options.append({
            "key": key,
            "text": text_en,
            "text_vi": text_vi
        })

    return options

def parse_answers(ans_raw):
    """
    Chuẩn hóa đáp án thành danh sách các chữ cái viết hoa (VD: 'A, C' -> ['A', 'C']).
    """
    if not ans_raw:
        return []
    return [ans.strip().upper() for ans in re.split(r'[,;\s]+', str(ans_raw).strip()) if ans.strip()]

def convert_single_excel(excel_path):
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    sheet = wb.active

    # Kiểm tra dòng tiêu đề (Header row)
    header_map = {}
    row_1_values = [sheet.cell(row=1, column=c).value for c in range(1, sheet.max_column + 1)]
    for idx, val in enumerate(row_1_values):
        if not val:
            continue
        v_clean = str(val).strip().lower()
        if any(k in v_clean for k in ["topic", "chủ đề", "domain"]):
            header_map["topic"] = idx + 1
        elif any(k in v_clean for k in ["question_vi", "câu hỏi vi", "dịch câu hỏi", "câu hỏi tiếng việt"]):
            header_map["question_vi"] = idx + 1
        elif any(k in v_clean for k in ["question", "câu hỏi"]):
            header_map["question"] = idx + 1
        elif any(k in v_clean for k in ["options_vi", "lựa chọn vi", "dịch lựa chọn", "dịch đáp án"]):
            header_map["options_vi"] = idx + 1
        elif any(k in v_clean for k in ["option", "lựa chọn", "các đáp án"]):
            header_map["options"] = idx + 1
        elif any(k in v_clean for k in ["answer", "đáp án", "correct"]):
            header_map["answer"] = idx + 1
        elif any(k in v_clean for k in ["explanation_vi", "giải thích vi", "dịch giải thích"]):
            header_map["explanation_vi"] = idx + 1
        elif any(k in v_clean for k in ["explanation", "giải thích"]):
            header_map["explanation"] = idx + 1

    # Nếu không nhận diện đủ các cột chính từ header, sử dụng quy tắc số cột mặc định
    use_named_headers = "question" in header_map and "options" in header_map

    questions = []
    start_row = 2 if use_named_headers or any(header_map.values()) else 1

    for r in range(start_row, sheet.max_row + 1):
        if use_named_headers:
            topic = sheet.cell(row=r, column=header_map.get("topic", 1)).value
            question_text = sheet.cell(row=r, column=header_map["question"]).value
            question_vi = sheet.cell(row=r, column=header_map["question_vi"]).value if "question_vi" in header_map else ""
            options_raw = sheet.cell(row=r, column=header_map["options"]).value
            options_vi_raw = sheet.cell(row=r, column=header_map["options_vi"]).value if "options_vi" in header_map else ""
            ans_raw = sheet.cell(row=r, column=header_map.get("answer", 4)).value
            explanation = sheet.cell(row=r, column=header_map["explanation"]).value if "explanation" in header_map else ""
            explanation_vi = sheet.cell(row=r, column=header_map["explanation_vi"]).value if "explanation_vi" in header_map else ""
        else:
            topic = sheet.cell(row=r, column=1).value
            question_text = sheet.cell(row=r, column=2).value
            options_raw = sheet.cell(row=r, column=3).value
            ans_raw = sheet.cell(row=r, column=4).value
            explanation = sheet.cell(row=r, column=5).value if sheet.max_column >= 5 else ""
            question_vi = sheet.cell(row=r, column=6).value if sheet.max_column >= 6 else ""
            options_vi_raw = sheet.cell(row=r, column=7).value if sheet.max_column >= 7 else ""
            explanation_vi = sheet.cell(row=r, column=8).value if sheet.max_column >= 8 else ""
        
        if not question_text or not str(question_text).strip():
            continue
            
        topic = str(topic).strip() if topic else "General"
        question_text = str(question_text).strip()
        question_vi = str(question_vi).strip() if question_vi else ""
        explanation = str(explanation).strip() if explanation else ""
        explanation_vi = str(explanation_vi).strip() if explanation_vi else ""
        
        parsed_options = parse_options(options_raw, options_vi_raw)
        parsed_answers = parse_answers(ans_raw)
        q_type = "multiple" if len(parsed_answers) > 1 else "single"
        
        questions.append({
            "id": len(questions) + 1,
            "topic": topic,
            "question": question_text,
            "question_vi": question_vi,
            "type": q_type,
            "options": parsed_options,
            "answer": parsed_answers,
            "explanation": explanation,
            "explanation_vi": explanation_vi
        })
    return questions

def extract_set_num(fname):
    """
    Trích xuất số thứ tự của bộ đề (VD: '1.xlsx' -> 1, 'Generative_AI_Leader_Full_Exam_60Q_Set10.xlsx' -> 10).
    """
    m = re.search(r'(?:Set|set|SET)[_\s]?(\d+)', fname)
    if m:
        return int(m.group(1))
    m2 = re.match(r'^(\d+)', fname)
    if m2:
        return int(m2.group(1))
    return None

def run_conversion():
    base_dir = Path(__file__).resolve().parent
    input_dir = base_dir / "input"
    output_dir = base_dir / "Quiz_data"
    os.makedirs(output_dir, exist_ok=True)
    
    excel_files = sorted(glob.glob(str(input_dir / "*.xlsx")))
    manifest_sets = []
    
    TOTAL_SETS = 10
    detected_sets = {}
    
    for fpath in excel_files:
        fname = Path(fpath).stem
        set_num = extract_set_num(fname)
        
        questions = convert_single_excel(fpath)
        set_id = f"set_{set_num}" if set_num else f"set_{fname}"
        
        # Ghi file json
        json_file = output_dir / f"{set_id}.json"
        with open(json_file, "w", encoding="utf-8") as f:
            json.dump(questions, f, ensure_ascii=False, indent=2)
            
        # Ghi file js (hỗ trợ offline và direct file://)
        js_file = output_dir / f"{set_id}.js"
        with open(js_file, "w", encoding="utf-8") as f:
            f.write(f"window.QUIZ_DATA_{set_id.upper()} = ")
            json.dump(questions, f, ensure_ascii=False, indent=2)
            f.write(";\n")
            
        detected_sets[set_num if set_num else fname] = {
            "id": set_id,
            "count": len(questions),
            "file": Path(fpath).name
        }
        print(f"✅ Đã xuất {set_id} ({len(questions)} câu hỏi) từ {Path(fpath).name}")

    # Tạo manifest cho 10 bộ đề
    for i in range(1, TOTAL_SETS + 1):
        if i in detected_sets:
            info = detected_sets[i]
            manifest_sets.append({
                "id": info["id"],
                "setNumber": i,
                "title": f"Bộ đề thi số {i}",
                "description": f"Bộ đề thực hành Generative AI Leader - {info['count']} câu hỏi trắc nghiệm chuẩn.",
                "questionCount": info["count"],
                "durationMinutes": 90,
                "isAvailable": True,
                "dataFile": f"{info['id']}.json",
                "jsFile": f"{info['id']}.js",
                "jsVar": f"QUIZ_DATA_{info['id'].upper()}"
            })
        else:
            manifest_sets.append({
                "id": f"set_{i}",
                "setNumber": i,
                "title": f"Bộ đề thi số {i}",
                "description": f"Đang chuẩn bị câu hỏi. Thêm file đề Set {i} vào input/ để kích hoạt.",
                "questionCount": 0,
                "durationMinutes": 90,
                "isAvailable": False,
                "dataFile": f"set_{i}.json",
                "jsFile": f"set_{i}.js",
                "jsVar": f"QUIZ_DATA_SET_{i}"
            })
            
    # Ghi manifest
    manifest_json = output_dir / "manifest.json"
    with open(manifest_json, "w", encoding="utf-8") as f:
        json.dump(manifest_sets, f, ensure_ascii=False, indent=2)

    manifest_js = output_dir / "manifest.js"
    with open(manifest_js, "w", encoding="utf-8") as f:
        f.write("window.QUIZ_MANIFEST = ")
        json.dump(manifest_sets, f, ensure_ascii=False, indent=2)
        f.write(";\n")
        
    # Đồng bộ questions.json mặc định trỏ về set_1
    if 1 in detected_sets:
        with open(output_dir / "questions.json", "w", encoding="utf-8") as f:
            with open(output_dir / "set_1.json", "r", encoding="utf-8") as src:
                f.write(src.read())

    print(f"\n🎉 HOÀN THÀNH CHUYỂN ĐỔI:")
    available_nums = [i for i in range(1, TOTAL_SETS + 1) if i in detected_sets]
    missing_nums = [i for i in range(1, TOTAL_SETS + 1) if i not in detected_sets]
    print(f"👉 Các đề đã sẵn sàng: {available_nums}")
    if missing_nums:
        print(f"👉 Các đề còn thiếu file trong input/: {missing_nums}")

if __name__ == "__main__":
    run_conversion()
